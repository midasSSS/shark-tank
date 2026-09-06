import csv
import hashlib
import io
from pathlib import Path

import fitz


MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_TOTAL_BYTES = 25 * 1024 * 1024


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract(name: str, data: bytes) -> tuple[list[dict], list[str]]:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"{name}: maximum file size is 15 MB.")
    extension = Path(name).suffix.lower()
    identifier = digest(data)[:12]
    sources, gaps = [], []

    def add(location: str, content: str):
        if not content.strip():
            gaps.append(f"{name} {location}: no readable content.")
            return
        # Every character is processed; splitting does not silently drop later pages.
        for offset in range(0, len(content), 12000):
            sources.append({"id": f"doc-{identifier}-{location}-{offset // 12000}", "title": name,
                            "location": location, "text": content[offset:offset + 12000],
                            "kind": "document", "hash": digest(data)})

    if extension == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            for number, page in enumerate(document, 1):
                text = page.get_text("text", sort=True)
                if len(text.strip()) < 40:
                    try:
                        text = page.get_text(textpage=page.get_textpage_ocr(full=True))
                    except Exception:
                        gaps.append(f"{name} page {number}: OCR unavailable; scanned content may be missing.")
                try:
                    tables = page.find_tables()
                    for table in tables.tables:
                        text += "\nTABLE\n" + "\n".join(" | ".join(str(cell or "") for cell in row) for row in table.extract())
                except Exception:
                    gaps.append(f"{name} page {number}: table extraction failed.")
                if page.get_images() and len(text.strip()) >= 40:
                    gaps.append(f"{name} page {number}: embedded images/charts require visual interpretation; text extraction alone does not verify them.")
                add(f"page-{number}", text)
    elif extension in {".csv", ".tsv"}:
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text), delimiter="\t" if extension == ".tsv" else ",")
        rows = list(reader)
        header = " | ".join(rows[0]) if rows else ""
        for start in range(1, len(rows), 100):
            add(f"rows-{start + 1}-{min(start + 100, len(rows))}", header + "\n" + "\n".join(" | ".join(row) for row in rows[start:start + 100]))
        if len(rows) <= 1:
            gaps.append(f"{name}: no data rows.")
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            for sheet, formula_sheet in zip(values, formulas):
                rows = []
                header = ""
                for n, (row, frow) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
                    for cell, fcell in zip(row, frow):
                        if fcell.data_type == "f" and cell.value is None:
                            gaps.append(f"{name} {sheet.title}!{fcell.coordinate}: formula has no cached value.")
                    line = f"{n}: " + " | ".join(str(cell.value if cell.value is not None else "") for cell in row)
                    if n == 1:
                        header = line
                    rows.append(line)
                    if len(rows) == 100:
                        add(f"{sheet.title}-through-row-{n}", header + "\n" + "\n".join(rows))
                        rows = []
                if rows:
                    add(f"{sheet.title}-final", header + "\n" + "\n".join(rows))
        finally:
            values.close()
            formulas.close()
    elif extension in {".txt", ".md"}:
        add("text", data.decode("utf-8-sig"))
    else:
        raise ValueError(f"Unsupported file type: {extension}")
    return sources, gaps
